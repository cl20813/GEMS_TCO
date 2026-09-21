from pathlib import Path
from shutil import copy2

from docx import Document


ROOT = Path("/Users/joonwonlee/Documents/GEMS_TCO-1")
SOURCE = ROOT / "career_materials/two_sigma/Joonwon_Lee_Two_Sigma_Quantitative_Researcher_Resume.docx"
OUT_DIR = ROOT / "career_materials/aquatic_capital"
OUTPUT = OUT_DIR / "Joonwon_Lee_Aquatic_Quantitative_Researcher_Resume.docx"


def set_body(paragraph, text):
    """Replace the text of a one-run body paragraph without changing its formatting."""
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def set_labeled_bullet(paragraph, label, body):
    """Preserve the template's bold label run and normal-weight body run."""
    paragraph.runs[0].text = label
    paragraph.runs[1].text = body
    for run in paragraph.runs[2:]:
        run.text = ""


def move_block_before(doc, start_index, end_index, before_index):
    """Move a contiguous paragraph block while preserving its complete OOXML formatting."""
    paragraphs = doc.paragraphs
    block = [paragraphs[i]._p for i in range(start_index, end_index + 1)]
    anchor = paragraphs[before_index]._p
    body = doc._body._element
    for element in block:
        body.remove(element)
    insertion_point = body.index(anchor)
    for offset, element in enumerate(block):
        body.insert(insertion_point + offset, element)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    set_body(
        doc.paragraphs[3],
        "Statistics Ph.D. candidate and computational statistician specializing in scalable likelihood inference, "
        "rigorous statistical diagnostics, and predictive modeling for stochastically dependent data. Designs and "
        "tests reproducible research systems in Python/PyTorch, integrates C++ for performance-critical components, "
        "and evaluates modeling assumptions through simulation and empirical evidence.",
    )

    doc.paragraphs[16].runs[0].text = "QUANTITATIVE RESEARCH"

    set_labeled_bullet(
        doc.paragraphs[18],
        "Scalable Inference: ",
        "Developed an advection-aware Vecchia likelihood approximation for nonseparable space-time covariance "
        "models, using ordered local conditional structure to estimate parameters and quantify uncertainty for up "
        "to 145,008 observations per day.",
    )
    set_labeled_bullet(
        doc.paragraphs[19],
        "Stochastic Diagnostics: ",
        "Developed scale- and frequency-resolved statistical diagnostics that identify where fitted stochastic "
        "models fail, separating persistent dependence from high-frequency noise, missingness, acquisition "
        "artifacts, and covariance misspecification.",
    )
    set_labeled_bullet(
        doc.paragraphs[20],
        "Data-Intensive Research Systems: ",
        "Built reproducible Python/PyTorch pipelines for quality filtering, irregular-to-regular spatial mapping, "
        "time-dependent alignment, optimization, simulation, and restartable CPU/GPU and HPC execution; integrated "
        "C++ ordering routines through pybind11.",
    )

    # Lead with the dissertation work because Aquatic is explicitly hiring completed
    # statistical/applied-mathematical researchers who can build research systems.
    move_block_before(doc, 16, 22, 10)

    doc.core_properties.title = "Joonwon Lee - Aquatic Capital Quantitative Researcher Resume"
    doc.core_properties.subject = "Application for Quantitative Researcher, PhD"
    doc.core_properties.keywords = "quantitative research, computational statistics, Python, C++, stochastic modeling"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
